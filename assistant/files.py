from __future__ import annotations

import base64
import io
import logging
from pathlib import PurePosixPath
from typing import Any

from telegram import Bot, Message

from .config import MAX_FILE_CHARS

log = logging.getLogger(__name__)

IMAGE_MEDIA_TYPES = {"image/jpeg", "image/png", "image/gif", "image/webp"}
TEXT_SUFFIXES = {
    ".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".yaml", ".yml", ".ini", ".toml",
    ".log", ".sql", ".sh", ".bat", ".py", ".js", ".ts", ".tsx", ".jsx", ".html", ".css",
    ".java", ".c", ".h", ".cpp", ".go", ".rs", ".php", ".rb", ".kt", ".swift", ".srt",
}


async def collect_blocks(message: Message, bot: Bot) -> tuple[list[dict[str, Any]], list[str]]:
    """Telegram xabaridagi matn va fayllarni model uchun content bloklarga aylantiradi."""
    blocks: list[dict[str, Any]] = []
    notes: list[str] = []

    if message.photo:
        data = await _download(message.photo[-1].file_id, bot)
        if data is None:
            notes.append("Couldn't download the photo (it might be over 20 MB).")
        else:
            blocks.append(_image_block(data, "image/jpeg"))

    document = message.document
    if document is not None:
        suffix = PurePosixPath(document.file_name or "").suffix.lower()
        mime = (document.mime_type or "").lower()
        data = await _download(document.file_id, bot)
        if data is None:
            notes.append(
                f"Couldn't download '{document.file_name}' "
                "(Telegram bots can only fetch files up to 20 MB)."
            )
        else:
            block, note = _document_block(data, document.file_name or "file", suffix, mime)
            if block is not None:
                blocks.append(block)
            if note:
                notes.append(note)

    if message.voice or message.audio or message.video or message.video_note:
        notes.append(
            "I can't understand voice or video messages yet — please send text instead."
        )

    text = message.text or message.caption
    if text:
        blocks.append({"type": "text", "text": text})

    return blocks, notes


async def _download(file_id: str, bot: Bot) -> bytes | None:
    try:
        tg_file = await bot.get_file(file_id)
        return bytes(await tg_file.download_as_bytearray())
    except Exception:
        log.exception("Faylni yuklab bo'lmadi: %s", file_id)
        return None


def _image_block(data: bytes, media_type: str) -> dict[str, Any]:
    encoded = base64.standard_b64encode(data).decode("ascii")
    return {"type": "image_url", "image_url": {"url": f"data:{media_type};base64,{encoded}"}}


def _document_block(
    data: bytes, name: str, suffix: str, mime: str
) -> tuple[dict[str, Any] | None, str]:
    if mime == "application/pdf" or suffix == ".pdf":
        text = _read_pdf(data)
        if not text.strip():
            return None, (
                f"Couldn't extract any text from '{name}' — it's probably a scanned "
                "(image-based) PDF, which I can't read."
            )
        return _text_block(name, text)

    if mime in IMAGE_MEDIA_TYPES:
        return _image_block(data, mime), ""

    if suffix == ".docx":
        return _text_block(name, _read_docx(data))

    if suffix in {".xlsx", ".xlsm"}:
        return _text_block(name, _read_xlsx(data))

    if suffix in TEXT_SUFFIXES or mime.startswith("text/") or mime == "application/json":
        return _text_block(name, data.decode("utf-8", errors="replace"))

    return None, f"I can't read this type of file yet: '{name}'."


def _text_block(name: str, content: str) -> tuple[dict[str, Any], str]:
    note = ""
    if len(content) > MAX_FILE_CHARS:
        content = content[:MAX_FILE_CHARS]
        note = f"'{name}' is too large — only the beginning was read."
        content += "\n\n[...file truncated...]"
    text = f"The user sent the file '{name}'. Contents:\n\n{content}"
    return {"type": "text", "text": text}, note


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _read_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_xlsx(data: bytes) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"## Varaq: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                parts.append(" | ".join(values))
    workbook.close()
    return "\n".join(parts)
